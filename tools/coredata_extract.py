#!/usr/bin/env python
"""xlsx → 중간 CSV. 코어 자료 반입의 앞단이다 (P17 4절).

**왜 두 단인가.** 뒷단(`ops/import_coredata.py`)은 규약대로 컨테이너 안에서
도는데, 거기서 xlsx 를 바로 읽으려면 웹 이미지에 `openpyxl` 이 들어가야 하고
NAS 공유(`/nfs/temp-share`)까지 컨테이너가 봐야 한다. 반입 하나 때문에 늘리기엔
넓다. 그리고 **1.1 의 단위 판단은 눈에 보이는 파일로 남아야 한다** — 그것이
`coredata/mapping.toml` 이고, 이 스크립트는 그 표만 따른다.

**Django 를 안 부르고 DB 도 안 만진다** — `ops/export_review.py`·
`ops/backup_db.py` 와 같은 자리라 호스트 venv 로 돈다.

    python tools/coredata_extract.py \
        --xlsx-dir /nfs/temp-share/DiaRUGA/coredata \
        --out      /data3/DiaRUGA/coredata

내는 것은 코어마다 둘이다:

    <지역>-<지점>.series.csv   key,label,unit,default_on,sort_order,origin
    <지역>-<지점>.points.csv   key,depth_mm,value

**깊이는 mm 정수로 낸다.** DB 가 그렇게 든다 — `(항목, 깊이)` 가 유일 제약의
열쇠라서 부동소수면 안 된다 (`CorePoint` 머리말).

**원본이 pptx 일 수도 있다** (206 · `RS21-diatom.pptx`). 엑셀에서 붙여 넣은
산점도가 값을 차트 캐시(`ppt/charts/chartN.xml` 의 `c:xVal`·`c:yVal`)에 그대로
들고 있어 원본 xlsx 없이도 읽힌다. 매핑표는 `[[core.block]]` 대신
`[[core.chart]]` 로 **몇 번째 슬라이드의 어느 제목 차트**인지, **깊이가 어느
축인지**를 적는다 — 같은 파일 안에서도 Opal 만 축이 뒤집혀 있어 짐작하면
안 되는 자리다. 뒤는 xlsx 와 같다(mm 정수 · 코어 길이 검사 · 같은 CSV).
"""
import argparse
import csv
import re
import sys
import tomllib
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

# **임포트에서 멈추지 않는다.** 시험이 `read_block`·`_to_mm` 을 xlsx 없이
# 부르는데(엑셀은 `iter_rows` 하나로만 쓴다), 여기서 `sys.exit` 하면 그 시험이
# openpyxl 을 깔아야만 도는 것이 된다 — 이 파일은 호스트 전용이고 웹 이미지에는
# 안 들어간다. 없으면 실제로 xlsx 를 열 때 말한다.
try:
    import openpyxl
except ImportError:                                          # pragma: no cover
    openpyxl = None

# 매핑표의 `expect_max_cm` 에서 이만큼 벗어나면 멈춘다. **단위를 잘못 읽으면
# 열 배로 어긋나므로** 넉넉히 잡아도 그 사고는 잡힌다. 코어 길이 자체가
# 어림값이라 좁게 잡으면 멀쩡한 반입이 멈춘다.
DEPTH_LO, DEPTH_HI = 0.5, 2.0


def _num(v):
    """숫자만 통과시킨다. `TYPE`·`Munsell C Hue` 같은 글자 칸은 여기서 걸린다."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).strip())
    except ValueError:
        return None


def _to_mm(depth: float, unit: str) -> int:
    """깊이를 mm 정수로. **반올림한다** — 0.1 cm 가 1 mm 다."""
    return round(depth * 10) if unit == "cm" else round(depth)


def read_block(ws, blk: dict) -> dict[str, dict[int, float]]:
    """블록 하나에서 `{key: {depth_mm: value}}` 를 뽑는다.

    **열은 번호로 짚는다.** `MS` 의 `Point`/`Whole`, `Spectro` 의 `SCE`/`SCI` 가
    같은 머리글을 반복해서 이름으로 짚으면 뒤엣것이 앞엣것을 덮는다.
    """
    dcol, dunit = blk["depth_col"], blk["depth_unit"]
    # 깊이 칸이 `GC03-C1 14` 처럼 코어 이름을 앞에 달고 있는 파일이 있다
    # (KPDC 의 `GC03-C1`, 197). **적힌 접두사만 뗀다** — 아무 글자나 벗기면
    # `LOD`·`Unit` 같은 머리 줄이 숫자로 읽힐 수 있다.
    prefix = blk.get("depth_prefix", "")
    cols = [(int(c[0]), c[1]) for c in blk["columns"]]
    out: dict[str, dict[int, float]] = {k: {} for _, k in cols}
    same, clash = 0, []
    for row in ws.iter_rows(min_row=blk["header_row"] + 1, values_only=True):
        raw = row[dcol - 1] if len(row) >= dcol else None
        if prefix and isinstance(raw, str):
            raw = raw[len(prefix):] if raw.startswith(prefix) else None
        depth = _num(raw)
        if depth is None:
            continue
        mm = _to_mm(depth, dunit)
        for col, key in cols:
            val = _num(row[col - 1]) if len(row) >= col else None
            # **값이 없는 깊이는 안 넣는다.** 넣으면 화면이 그 구간을 이어
            # 그려 안 잰 구간이 잰 것처럼 뜬다.
            if val is None:
                continue
            if mm in out[key]:
                # **같은 값이면 붙여넣기가 겹친 것이고, 다르면 사람이 정할
                # 일이다.** `RS14-GC04` 의 `MS` 가 221~260 cm 40점을 한 번 더
                # 들고 있는데 값이 글자 그대로 같다 — 그건 버려도 된다. 값이
                # 다르면 어느 쪽이 맞는지 이 스크립트가 고를 수 없다.
                if out[key][mm] != val:
                    clash.append((key, mm, out[key][mm], val))
                else:
                    same += 1
                continue
            out[key][mm] = val
    if same:
        print(f"      같은 깊이·같은 값이 {same}개 겹쳐 있어 한 번만 넣습니다 "
              f"(시트 {blk['sheet']})")
    if clash:
        print(f"      !! 같은 깊이에 다른 값이 있습니다 (시트 {blk['sheet']}, "
              f"{len(clash)}건) — 어느 쪽이 맞는지 사람이 정해야 합니다")
        for key, mm, a, b in clash[:5]:
            print(f"         {key} {mm / 10:g} cm: {a:g} vs {b:g}")
        raise ValueError(f"{blk['sheet']}: 같은 깊이에 다른 값")
    return out


_NS_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_NS_C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"


def _norm_title(t: str) -> str:
    """차트 제목을 비교용으로. `opal  final`(공백 둘)·`Density(g/cm3)` 처럼
    같은 파일 안에서도 띄어쓰기·대소문자가 들쭉날쭉하다."""
    return " ".join(t.lower().split())


def read_pptx_charts(path: Path) -> dict[int, dict[str, tuple[dict, dict]]]:
    """pptx 의 차트를 `{슬라이드 번호: {제목: ({idx: x}, {idx: y})}}` 로 읽는다.

    산점도 하나에 계열 하나인 파일만 상정한다 — 계열이 둘이면 어느 것인지
    매핑표가 말할 자리가 없으므로 멈춘다. 같은 슬라이드에 같은 제목이 둘이어도
    같은 이유로 멈춘다.
    """
    z = zipfile.ZipFile(path)
    out: dict[int, dict[str, tuple[dict, dict]]] = {}
    for name in z.namelist():
        m = re.fullmatch(r"ppt/slides/slide(\d+)\.xml", name)
        if not m:
            continue
        sno = int(m.group(1))
        rels = z.read(f"ppt/slides/_rels/slide{sno}.xml.rels").decode()
        charts = out.setdefault(sno, {})
        for cm in re.finditer(r'Target="\.\./charts/(chart\d+\.xml)"', rels):
            root = ET.fromstring(z.read("ppt/charts/" + cm.group(1)))
            title = _norm_title(" ".join(t.text or "" for t in root.iter(_NS_A + "t")))
            sers = list(root.iter(_NS_C + "ser"))
            if len(sers) != 1:
                raise ValueError(f"슬라이드 {sno} '{title}': 계열이 {len(sers)}개")
            if title in charts:
                raise ValueError(f"슬라이드 {sno}: 제목 '{title}' 이 둘")
            # **점은 `idx` 로 짝짓는다.** 셀이 비면 그 축의 캐시에서 그 점만
            # 빠지므로 두 축을 차례대로 zip 하면 그 뒤가 전부 한 칸씩 어긋난다.
            charts[title] = tuple(
                {int(pt.get("idx")): _num(pt.findtext(_NS_C + "v"))
                 for pt in sers[0].findall(f".//{_NS_C}{ax}//{_NS_C}pt")}
                for ax in ("xVal", "yVal"))
    return out


def read_chart(charts: dict, ch: dict) -> dict[int, float]:
    """`[[core.chart]]` 하나 → `{depth_mm: value}`.

    **깊이 축은 매핑표가 말한다**(`depth_axis` = `x`|`y`). 그래도 읽은 깊이가
    0 이상이고 내려가는 차례인지는 본다 — 축을 반대로 적으면 값이 깊이 자리에
    앉는데, 계측값은 오르내리므로 여기서 걸린다.
    """
    sno, title = int(ch["slide"]), _norm_title(ch["title"])
    if sno not in charts or title not in charts[sno]:
        have = ", ".join(sorted(charts.get(sno, {})))
        raise ValueError(f"슬라이드 {sno} 에 '{title}' 차트가 없습니다 (있는 것: {have})")
    xs, ys = charts[sno][title]
    depths, vals = (ys, xs) if ch["depth_axis"] == "y" else (xs, ys)
    out: dict[int, float] = {}
    prev = -1
    for i in sorted(depths):
        d, v = depths[i], vals.get(i)
        # 값이 없는 깊이는 안 넣는다 — `read_block` 과 같은 이유.
        if d is None or v is None:
            continue
        if d < 0 or d < prev:
            raise ValueError(f"슬라이드 {sno} '{title}': 깊이가 {prev:g} 다음에 "
                             f"{d:g} — depth_axis 가 반대로 적혀 있지 않은지")
        prev = d
        mm = _to_mm(d, ch["depth_unit"])
        if mm in out and out[mm] != v:
            raise ValueError(f"슬라이드 {sno} '{title}': {d:g} cm 에 값이 둘")
        out[mm] = v
    return out


def _entries(core: dict) -> list[tuple[dict, list[tuple[str, str, str]]]]:
    """블록이든 차트든 `(설정, [(key, label, unit)…])` 로 편다."""
    if "chart" in core:
        return [(ch, [(ch["key"], ch["label"], ch["unit"])]) for ch in core["chart"]]
    return [(b, [(c[1], c[2], c[3]) for c in b["columns"]]) for b in core["block"]]


def extract(core: dict, xlsx_dir: Path, out_dir: Path) -> int:
    name = f"{core['site']}-{core['locality']}"
    src = xlsx_dir / core["file"]
    if not src.exists():
        print(f"  !! 원본이 없습니다: {src}")
        return 1
    if openpyxl is None and src.suffix.lower() != ".pptx":
        print("  !! openpyxl 이 없습니다 — pip install openpyxl")
        return 1
    print(f"  {name}  ← {core['file']}")
    is_pptx = src.suffix.lower() == ".pptx"
    if is_pptx:
        try:
            charts = read_pptx_charts(src)
        except ValueError as e:
            print(f"      !! {e}")
            return 1
        wb = None
    else:
        wb = openpyxl.load_workbook(src, read_only=True, data_only=True)

    default_on = set(core.get("default_on", []))
    meta: dict[str, dict] = {}
    points: dict[str, dict[int, float]] = {}
    for blk, cols in _entries(core):
        if is_pptx:
            try:
                got = {cols[0][0]: read_chart(charts, blk)}
            except ValueError as e:
                print(f"      !! {e}")
                return 1
            where = f"slide{blk['slide']}:{_norm_title(blk['title'])}"
        else:
            if blk["sheet"] not in wb.sheetnames:
                print(f"      !! 시트가 없습니다: {blk['sheet']}")
                return 1
            try:
                got = read_block(wb[blk["sheet"]], blk)
            except ValueError:
                # `read_block` 이 이미 무엇이 어긋났는지 적었다. 여기서 스택을
                # 쏟으면 그 줄이 묻힌다.
                wb.close()
                return 1
            where = blk["sheet"]
        for i, (key, label, unit) in enumerate(cols):
            if key in meta:
                print(f"      !! key 가 겹칩니다: {key}")
                return 1
            meta[key] = {
                "key": key, "label": label, "unit": unit,
                "default_on": 1 if key in default_on else 0,
                "sort_order": blk.get("sort_order", 100) + i,
                "origin": f"{core['file']}::{where}",
            }
            points[key] = got[key]
    if wb is not None:
        wb.close()

    # **사람이 적어 둔 코어 길이로 자기를 검사한다.** 단위를 잘못 읽으면 열
    # 배로 어긋나므로 여기서 걸린다 — 매핑표가 짐작을 막는 것과 짝이다.
    expect = core.get("expect_max_cm")
    bad = []
    if expect:
        for key, pts in points.items():
            if not pts:
                continue
            deep = max(pts) / 10
            if not (expect * DEPTH_LO <= deep <= expect * DEPTH_HI):
                bad.append((key, deep))
    if bad:
        print(f"      !! 코어 길이가 {expect} cm 인데 깊이가 벗어납니다 — "
              f"매핑표의 depth_unit 을 보세요")
        for key, deep in bad:
            print(f"         {key}: 최대 {deep:g} cm")
        return 1

    # 켜라고 적어 둔 항목이 실제로 있는가. 오타면 화면이 아무것도 안 켠 채
    # 뜨는데, 그것은 예외도 경고도 없는 종류의 고장이다.
    for key in sorted(default_on - set(meta)):
        print(f"      !! default_on 에 없는 항목이 적혀 있습니다: {key}")
        return 1

    out_dir.mkdir(parents=True, exist_ok=True)
    with (out_dir / f"{name}.series.csv").open("w", newline="") as f:
        w = csv.DictWriter(f, ["key", "label", "unit", "default_on",
                               "sort_order", "origin"])
        w.writeheader()
        for key in sorted(meta, key=lambda k: meta[k]["sort_order"]):
            w.writerow(meta[key])
    n = 0
    with (out_dir / f"{name}.points.csv").open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["key", "depth_mm", "value"])
        for key in sorted(points, key=lambda k: meta[k]["sort_order"]):
            for mm in sorted(points[key]):
                w.writerow([key, mm, repr(points[key][mm])])
                n += 1
    empty = [k for k, v in points.items() if not v]
    print(f"      항목 {len(meta)}개 · 점 {n:,}개 → {out_dir / (name + '.*.csv')}")
    if empty:
        print(f"      점이 하나도 없는 항목 {len(empty)}개: {', '.join(sorted(empty))}")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description="코어 자료 xlsx → 중간 CSV")
    ap.add_argument("--mapping", default=str(Path(__file__).resolve().parent.parent
                                             / "coredata" / "mapping.toml"))
    ap.add_argument("--xlsx-dir", required=True)
    ap.add_argument("--out", required=True,
                    help="컨테이너도 보는 자리여야 한다 (/data3/DiaRUGA/coredata)")
    ap.add_argument("--only", default="", help="지역-지점 하나만 (RS14-GC04)")
    a = ap.parse_args()

    with open(a.mapping, "rb") as f:
        conf = tomllib.load(f)
    print(f"매핑표: {a.mapping}")
    rc = 0
    for core in conf["core"]:
        name = f"{core['site']}-{core['locality']}"
        if a.only and a.only != name:
            continue
        rc |= extract(core, Path(a.xlsx_dir), Path(a.out))
    return rc


if __name__ == "__main__":
    sys.exit(main())
